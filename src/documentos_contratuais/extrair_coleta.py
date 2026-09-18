"""Extrai os campos de um documento jurídico a partir do .docx "Coleta de
Dados" preenchido pelo cliente (§ Contratos).

⭐ 2026-09-18 — porta de `contratos-backend/src/documentos/extrair_coleta.py`,
trocando só `_dados_iniciais_contrato()` (Contratos) por
`dados_iniciais_contrato()` (`dados_documento_contratual.py`, ATLAS) — a
lógica de extração em si não muda uma linha.

Formato do arquivo: uma tabela principal de 2 colunas (rótulo | valor) com
linhas de cabeçalho de seção mescladas (célula 0 == célula 1: "INFORMAÇÕES
DA CLIENTE" / "REPRESENTANTE LEGAL" / "INFORMAÇÕES SOBRE O SERVIÇO" / "CUSTO
DO PROJETO"), mais uma tabela "TESTEMUNHA 1" e uma "TESTEMUNHA 2" separadas,
no mesmo formato rótulo|valor.

Regra central: os valores são texto livre digitado por gente — rótulos têm
espaço sobrando e negrito inconsistente (por isso o matching é por rótulo
NORMALIZADO), e os valores só viram dado estruturado quando dá pra reconhecer
o formato (CPF/CNPJ/telefone/moeda/data). O que não dá, vira pendência —
nunca inventamos valor. Um rótulo ausente ou uma tabela fora do padrão também
vira pendência; nunca derruba a extração inteira.
"""

import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from typing import Any, Dict, List, Optional, Tuple

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from src.utils.dados_documento_contratual import dados_iniciais_contrato

MESES_ORDEM = [
    "janeiro", "fevereiro", "março", "abril", "maio", "junho",
    "julho", "agosto", "setembro", "outubro", "novembro", "dezembro",
]


def normalizar_rotulo(texto: str) -> str:
    """strip + lower + sem acento + espaços colapsados + sem pontuação de borda
    (":" / "?") — pra casar rótulo digitado com espaço sobrando/maiúscula
    inconsistente ("CNPJ " vs "cnpj", "Será parcelado ?" vs "sera parcelado")."""
    texto = unicodedata.normalize("NFKD", texto.strip())
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = re.sub(r"\s+", " ", texto.lower())
    return texto.strip(" ?:")


def _eh_opcional(valor_bruto: str) -> bool:
    # normalizar_rotulo não tira parênteses (é usado em rótulos também, onde
    # eles fazem parte do texto, ex. "RG (com órgão emissor)") — aqui, no
    # valor, "(Opcional)" tem parênteses de propósito.
    return normalizar_rotulo(valor_bruto).strip("()") == "opcional"


def _somente_digitos(texto: str) -> str:
    return re.sub(r"\D", "", texto)


def validar_cpf(digitos: str) -> bool:
    if len(digitos) != 11 or len(set(digitos)) == 1:
        return False
    nums = [int(d) for d in digitos]

    def dv(base: List[int], inicio: int) -> int:
        s = sum(n * f for n, f in zip(base, range(inicio, 1, -1)))
        r = s % 11
        return 0 if r < 2 else 11 - r

    d1 = dv(nums[:9], 10)
    d2 = dv(nums[:9] + [d1], 11)
    return nums[9] == d1 and nums[10] == d2


def validar_cnpj(digitos: str) -> bool:
    if len(digitos) != 14 or len(set(digitos)) == 1:
        return False
    nums = [int(d) for d in digitos]
    w1 = [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]
    w2 = [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]

    def dv(base: List[int], pesos: List[int]) -> int:
        s = sum(n * p for n, p in zip(base, pesos))
        r = s % 11
        return 0 if r < 2 else 11 - r

    d1 = dv(nums[:12], w1)
    d2 = dv(nums[:12] + [d1], w2)
    return nums[12] == d1 and nums[13] == d2


def formatar_cpf(digitos: str) -> str:
    d = digitos[:11]
    r = d[:3]
    if len(d) > 3:
        r += "." + d[3:6]
    if len(d) > 6:
        r += "." + d[6:9]
    if len(d) > 9:
        r += "-" + d[9:11]
    return r


def formatar_cnpj(digitos: str) -> str:
    d = digitos[:14]
    r = d[:2]
    if len(d) > 2:
        r += "." + d[2:5]
    if len(d) > 5:
        r += "." + d[5:8]
    if len(d) > 8:
        r += "/" + d[8:12]
    if len(d) > 12:
        r += "-" + d[12:14]
    return r


# Caracteres que ainda fazem parte do NÚMERO do RG (dígitos, o verificador
# "X" e a pontuação usual). O primeiro caractere fora disso — tipicamente a
# primeira letra do órgão emissor ("SSP/SC") — encerra o número.
_RG_CHARS_NUMERO = set("0123456789X.- ")


def separar_rg(texto: str) -> Tuple[str, str]:
    """"2027163 SSP/SC" -> ("2027163", "SSP/SC"). O campo é rotulado "RG (com
    órgão emissor)", então o valor quase sempre traz as duas coisas juntas —
    e o órgão precisa sobreviver inteiro (vai impresso no documento, em "RG
    nº. ..."). Quando o RG começa por letra ("MG-12.345.678"), o número sai
    vazio e o valor todo fica na segunda parte, preservado como veio."""
    upper = texto.strip().upper()
    i = 0
    while i < len(upper) and upper[i] in _RG_CHARS_NUMERO:
        i += 1
    return upper[:i].strip(), upper[i:].strip()


def formatar_rg(texto: str) -> str:
    """Formata o número no agrupamento 2.3.3-1 (o mais usual, de SP) SÓ quando
    ele tem 8 ou 9 caracteres. RG não é padronizado nacionalmente e outros
    estados emitem números mais curtos ou mais longos (SC com 7 dígitos, RJ com
    10) — nesses casos o número fica exatamente como o cliente escreveu, em vez
    de ganhar pontos em posição inventada ("2027163" -> "20.271.63"). O órgão
    emissor é sempre preservado no fim."""
    numero, orgao = separar_rg(texto)
    chars = re.sub(r"[^0-9X]", "", numero)
    if len(chars) in (8, 9):
        r = f"{chars[:2]}.{chars[2:5]}.{chars[5:8]}"
        if len(chars) == 9:
            r += f"-{chars[8]}"
    else:
        r = re.sub(r"\s+", " ", numero)
    return f"{r} {orgao}".strip()


def formatar_telefone(digitos: str) -> str:
    d = digitos[:11]
    ddd, resto = d[:2], d[2:]
    r = f"({ddd}) " if len(ddd) == 2 else (f"({ddd}" if ddd else "")
    if len(d) <= 10:
        r += f"{resto[:4]}-{resto[4:8]}" if len(resto) > 4 else resto
    else:
        r += f"{resto[:5]}-{resto[5:9]}" if len(resto) > 5 else resto
    return r


def juntar_linhas(texto: str) -> str:
    """Junta num valor de uma linha só um texto que veio quebrado em várias
    linhas/tabs na célula da coleta. O separador esperado (do endereço, p.ex.)
    é a vírgula; quando o cliente usa Enter ou Tab entre logradouro/bairro/
    cidade, a célula vem com "\\n"/"\\t" no meio — aqui trocamos essas quebras
    por ", " pra sair uma linha contínua (sem duplicar vírgula onde o cliente
    já tinha posto uma no fim do trecho)."""
    partes = [p.strip(" ,") for p in re.split(r"[\n\r\t]+", texto)]
    return ", ".join(p for p in partes if p)


_CONECTIVOS = {"de", "da", "do", "das", "dos", "e"}


def capitalizar_nome(texto: str) -> str:
    palavras = texto.split(" ")
    resultado = []
    for i, palavra in enumerate(palavras):
        if not palavra:
            resultado.append(palavra)
            continue
        minuscula = palavra.lower()
        resultado.append(minuscula if i > 0 and minuscula in _CONECTIVOS else minuscula[0].upper() + minuscula[1:])
    return " ".join(resultado)


def parse_moeda(texto: str) -> Optional[float]:
    """"R$ 22.000,00" -> 22000.0. None se não achar nenhum número no texto."""
    m = re.search(r"(\d{1,3}(?:\.\d{3})*|\d+)(,\d{1,2})?", texto)
    if not m:
        return None
    inteiro = m.group(1).replace(".", "")
    centavos = (m.group(2) or ",00")[1:].ljust(2, "0")
    return int(inteiro) + int(centavos) / 100


def parse_inteiro(texto: str) -> Optional[int]:
    m = re.search(r"\d+", texto)
    return int(m.group(0)) if m else None


def parse_booleano(texto: str) -> Optional[bool]:
    n = normalizar_rotulo(texto)
    if n == "sim":
        return True
    if n in ("nao", "não"):
        return False
    return None


_RE_DATA_BR = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{2,4})$")
_RE_DATA_ISO = re.compile(r"^(\d{4})-(\d{2})-(\d{2})$")


def parse_data_iso(texto: str) -> Optional[str]:
    """Só reconhece formato de data fechado (DD/MM/AAAA ou AAAA-MM-DD) — texto
    como "30/07 e todos os dias 30..." não bate e vira pendência, de propósito:
    é descrição de recorrência, não uma data única."""
    texto = texto.strip()
    m = _RE_DATA_ISO.match(texto)
    if m:
        ano, mes, dia = m.groups()
        return f"{ano}-{mes}-{dia}"
    m = _RE_DATA_BR.match(texto)
    if m:
        dia, mes, ano = m.groups()
        if len(ano) == 2:
            ano = f"20{ano}"
        try:
            dia_i, mes_i = int(dia), int(mes)
        except ValueError:
            return None
        if not (1 <= dia_i <= 31 and 1 <= mes_i <= 12):
            return None
        return f"{int(ano):04d}-{mes_i:02d}-{dia_i:02d}"
    return None


_RE_ETAPA = re.compile(r"^(?P<nome>.*?)\s*\((?P<dias>\d+)\s*dias?\s*[uú]teis?\)\s*$", re.IGNORECASE)


def extrair_etapas(texto: str) -> Tuple[List[dict], List[str]]:
    """"Ambientação (5 dias úteis) + Análise Mercadológica (30 dias úteis)" ->
    lista de escopos. Cada trecho separado por "+" que não bater no formato
    "Nome (N dias úteis)" vira pendência — nunca inventa o prazo."""
    escopos: List[dict] = []
    pendencias: List[str] = []
    for trecho in texto.split("+"):
        trecho = trecho.strip()
        if not trecho:
            continue
        m = _RE_ETAPA.match(trecho)
        if m:
            escopos.append({"nome": m.group("nome").strip(), "prazo_dias_uteis": int(m.group("dias"))})
        else:
            pendencias.append(f'Etapa do serviço não reconhecida automaticamente: "{trecho}" — adicione manualmente.')
    return escopos, pendencias


@dataclass(frozen=True)
class CampoColeta:
    caminho: str
    tipo: str
    rotulo: str


# Caminho usado nas pendências que não são de um campo específico (tabela do
# documento fora do padrão). Nunca é filtrado por tipo de documento: vale pra
# quem estiver aproveitando a coleta, seja qual for o documento.
CAMPO_DOCUMENTO = "documento"


@dataclass(frozen=True)
class PendenciaColeta:
    """Uma pendência de extração + o caminho do campo (no formato de `dados`
    do Contrato de PS) a que ela se refere. O caminho existe porque a Coleta
    é aproveitada por documentos que só usam parte dela — o NDA não tem
    financeiro, então uma pendência de valor de parcela não é assunto dele
    (ver `adaptar_coleta_para_tipo.py`)."""

    campo: str
    mensagem: str


# tipo -> como normalizar o valor bruto da célula. Cada um devolve
# (valor_normalizado, pendencia_ou_none).
def _normalizar_valor(tipo: str, bruto: str, rotulo: str) -> Tuple[Any, Optional[str]]:
    valor = bruto.strip()
    if _eh_opcional(valor):
        # Critério dado: "(Opcional)" como valor = campo vazio, sem pendência.
        vazio_por_tipo: Dict[str, Any] = {"inteiro": 0, "moeda": 0.0, "booleano": False}
        return vazio_por_tipo.get(tipo, ""), None
    if not valor:
        return ("" if tipo not in ("inteiro", "moeda", "booleano") else vazio_padrao(tipo)), None

    if tipo == "texto":
        return juntar_linhas(valor), None
    if tipo == "texto_upper":
        return juntar_linhas(valor).upper(), None
    if tipo == "nome":
        return capitalizar_nome(juntar_linhas(valor)), None
    if tipo == "cpf":
        d = _somente_digitos(valor)
        if len(d) != 11:
            return valor, f'{rotulo} não reconhecido automaticamente ("{valor}") — confira e corrija.'
        return formatar_cpf(d), None
    if tipo == "cnpj":
        d = _somente_digitos(valor)
        if len(d) != 14:
            return valor, f'{rotulo} não reconhecido automaticamente ("{valor}") — confira e corrija.'
        return formatar_cnpj(d), None
    if tipo == "rg":
        # Ao contrário de CPF/CNPJ, aqui não dá pra exigir um tamanho: cada
        # estado emite RG com uma quantidade de dígitos diferente (7 em SC, 8-9
        # em SP, 10 no RJ). Só vira pendência quando não há número nenhum
        # reconhecível — 5 dígitos é um piso folgado, abaixo disso é texto
        # ("não sei", "vou mandar depois"), não um documento.
        if len(_somente_digitos(valor)) < 5:
            return valor, f'{rotulo} não reconhecido automaticamente ("{valor}") — confira e corrija.'
        return formatar_rg(valor), None
    if tipo == "telefone":
        d = _somente_digitos(valor)
        if len(d) not in (10, 11):
            return valor, f'{rotulo} não reconhecido automaticamente ("{valor}") — confira e corrija.'
        return formatar_telefone(d), None
    if tipo == "moeda":
        m = parse_moeda(valor)
        if m is None:
            return 0.0, f'{rotulo} não reconhecido automaticamente ("{valor}") — preencha manualmente.'
        return m, None
    if tipo == "inteiro":
        i = parse_inteiro(valor)
        if i is None:
            return 0, f'{rotulo} não reconhecido automaticamente ("{valor}") — preencha manualmente.'
        return i, None
    if tipo == "booleano":
        b = parse_booleano(valor)
        if b is None:
            return False, f'{rotulo} não reconhecido automaticamente ("{valor}") — confira e corrija.'
        return b, None
    if tipo == "data_iso":
        data = parse_data_iso(valor)
        if data is None:
            return "", f'{rotulo} não reconhecido automaticamente ("{valor}") — preencha manualmente.'
        return data, None
    return valor, None


def vazio_padrao(tipo: str) -> Any:
    return {"inteiro": 0, "moeda": 0.0, "booleano": False}.get(tipo, "")


CAMPOS_CLIENTE = {
    "razao social": CampoColeta("contratante.razao_social", "texto_upper", "Razão social"),
    "cnpj": CampoColeta("contratante.cnpj", "cnpj", "CNPJ do contratante"),
    "endereco completo [logradouro, cep, bairro, cidade, estado]": CampoColeta(
        "contratante.endereco", "texto", "Endereço do contratante"
    ),
}
CAMPOS_REPRESENTANTE = {
    "email": CampoColeta("contratante.representante.email", "texto", "E-mail do representante"),
    "telefone": CampoColeta("contratante.representante.telefone", "telefone", "Telefone do representante"),
    "cargo do representante legal": CampoColeta("contratante.representante.cargo", "texto", "Cargo do representante"),
    "nome do representante legal": CampoColeta("contratante.representante.nome", "nome", "Nome do representante"),
    "nacionalidade": CampoColeta("contratante.representante.nacionalidade", "texto", "Nacionalidade do representante"),
    "estado civil": CampoColeta("contratante.representante.estado_civil", "texto", "Estado civil do representante"),
    "profissao": CampoColeta("contratante.representante.profissao", "texto", "Profissão do representante"),
    "cpf do representante legal": CampoColeta("contratante.representante.cpf", "cpf", "CPF do representante"),
    "rg do representante (com orgao emissor)": CampoColeta(
        "contratante.representante.rg", "rg", "RG do representante"
    ),
    "endereco completo [logradouro, cep, bairro, cidade, estado]": CampoColeta(
        "contratante.representante.endereco", "texto", "Endereço do representante"
    ),
}
CAMPOS_SERVICO = {
    "qual servico sera prestado": CampoColeta("projeto.servico", "texto", "Serviço a ser prestado"),
    "quantos consultores serao disponibilizados": CampoColeta(
        "projeto.num_consultores", "inteiro", "Número de consultores"
    ),
    "quantos coordenadores serao disponibilizados": CampoColeta(
        "projeto.num_coordenadores", "inteiro", "Número de coordenadores"
    ),
    "data de inicio do servico": CampoColeta("projeto.data_inicio", "texto", "Data de início do serviço"),
    "data de termino do servico": CampoColeta("projeto.data_termino", "texto", "Data de término do serviço"),
}
# Tratado à parte (vira lista de escopos, não um campo escalar).
ROTULO_ETAPAS = "quais as etapas do servico (com o prazo estipulado para cada etapa)"

CAMPOS_CUSTO = {
    "qual o valor do projeto": CampoColeta("financeiro.valor_total", "moeda", "Valor do projeto"),
    "sera parcelado": CampoColeta("financeiro.parcelado", "booleano", "Parcelamento"),
    "quantas parcelas serao": CampoColeta("financeiro.numero_parcelas", "inteiro", "Número de parcelas"),
    "quando ira comecar o pagamento": CampoColeta(
        "financeiro.primeiro_vencimento", "data_iso", "Data do primeiro vencimento"
    ),
    "em que dia do mes o pagamento sera realizado": CampoColeta(
        "financeiro.dia_vencimento_mensal", "inteiro", "Dia do vencimento mensal"
    ),
    "qual sera a forma de pagamento (pix, boleto)": CampoColeta(
        "financeiro.forma_pagamento", "texto_upper", "Forma de pagamento"
    ),
    # "Qual o valor das parcelas" fica de fora de propósito: é sempre derivado
    # (valor_total / numero_parcelas), igual ao formulário manual — nunca lido
    # do documento, pra não divergir de uma conta que o cliente errou.
}

SECOES_CAMPOS: Dict[str, Dict[str, CampoColeta]] = {
    "informacoes da cliente": CAMPOS_CLIENTE,
    "representante legal": CAMPOS_REPRESENTANTE,
    "informacoes sobre o servico": CAMPOS_SERVICO,
    "custo do projeto": CAMPOS_CUSTO,
}

CAMPOS_TESTEMUNHA = {
    "nome": CampoColeta("nome", "nome", "Nome"),
    "cpf": CampoColeta("cpf", "cpf", "CPF"),
}


def _set_path(dados: dict, caminho: str, valor: Any) -> None:
    partes = caminho.split(".")
    alvo = dados
    for parte in partes[:-1]:
        alvo = alvo[parte]
    alvo[partes[-1]] = valor


def _linha_e_cabecalho_secao(celulas: List[str]) -> bool:
    texto = celulas[0].strip()
    return bool(texto) and texto == celulas[1].strip()


def _processar_tabela_principal(
    tabela, dados: dict, pendencias: List[PendenciaColeta], encontrados: set
) -> None:
    secao_atual: Optional[str] = None
    for linha in tabela.rows:
        celulas = [c.text for c in linha.cells]
        if _linha_e_cabecalho_secao(celulas):
            rotulo_secao = normalizar_rotulo(celulas[0])
            if rotulo_secao in SECOES_CAMPOS:
                secao_atual = rotulo_secao
            continue
        rotulo = normalizar_rotulo(celulas[0])
        if not rotulo or secao_atual is None:
            continue

        if secao_atual == "informacoes sobre o servico" and rotulo == ROTULO_ETAPAS:
            encontrados.add((secao_atual, rotulo))
            escopos, pendencias_etapas = extrair_etapas(celulas[1])
            dados["projeto"]["escopos"] = escopos
            pendencias.extend(PendenciaColeta("projeto.escopos", m) for m in pendencias_etapas)
            continue

        campo = SECOES_CAMPOS[secao_atual].get(rotulo)
        if campo is None:
            continue
        encontrados.add((secao_atual, rotulo))
        valor, pendencia = _normalizar_valor(campo.tipo, celulas[1], campo.rotulo)
        _set_path(dados, campo.caminho, valor)
        if pendencia:
            pendencias.append(PendenciaColeta(campo.caminho, pendencia))


def _processar_tabela_testemunha(tabela, indice: int, dados: dict, pendencias: List[PendenciaColeta]) -> None:
    encontrados: set = set()
    for linha in tabela.rows[1:]:
        celulas = [c.text for c in linha.cells]
        rotulo = normalizar_rotulo(celulas[0])
        campo = CAMPOS_TESTEMUNHA.get(rotulo)
        if campo is None:
            continue
        encontrados.add(rotulo)
        valor, pendencia = _normalizar_valor(campo.tipo, celulas[1], f"{campo.rotulo} da Testemunha {indice + 1}")
        dados["testemunhas"][indice][campo.caminho] = valor
        if pendencia:
            pendencias.append(PendenciaColeta(f"testemunhas.{indice}.{campo.caminho}", pendencia))
    for rotulo, campo in CAMPOS_TESTEMUNHA.items():
        if rotulo not in encontrados:
            pendencias.append(
                PendenciaColeta(
                    f"testemunhas.{indice}.{campo.caminho}",
                    f"Testemunha {indice + 1}: \"{campo.rotulo}\" não encontrado no documento.",
                )
            )


_RE_TESTEMUNHA_HEADER = re.compile(r"^testemunha (\d+)$")


def extrair_dados_coleta(conteudo: bytes) -> Tuple[dict, List[str]]:
    """Ponto de entrada: bytes do .docx de Coleta de Dados -> (dados no
    formato de `dados` do Contrato de PS, pendências de extração em português).

    Levanta ValueError se o arquivo não for um .docx válido — o chamador
    traduz isso num erro amigável (`RegraDeNegocioError`)."""
    dados, pendencias = extrair_dados_coleta_detalhado(conteudo)
    return dados, [p.mensagem for p in pendencias]


def extrair_dados_coleta_detalhado(conteudo: bytes) -> Tuple[dict, List[PendenciaColeta]]:
    """Mesma extração de `extrair_dados_coleta`, com cada pendência
    acompanhada do campo a que se refere — pra quem precisa filtrar as
    pendências pelo que o documento de destino de fato aproveita da Coleta
    (ver `adaptar_coleta_para_tipo.py`)."""
    try:
        doc = Document(io.BytesIO(conteudo))
    except (PackageNotFoundError, zipfile.BadZipFile):
        raise ValueError("O arquivo enviado não é um .docx válido.")

    dados = dados_iniciais_contrato()
    pendencias: List[PendenciaColeta] = []
    encontrados: set = set()

    if not doc.tables:
        raise ValueError("O documento não tem nenhuma tabela — não parece ser uma Coleta de Dados válida.")

    for tabela in doc.tables:
        if not tabela.rows:
            continue
        primeira = [c.text for c in tabela.rows[0].cells]
        if not _linha_e_cabecalho_secao(primeira):
            pendencias.append(
                PendenciaColeta(
                    CAMPO_DOCUMENTO,
                    f'Não foi possível interpretar uma tabela do documento (começa com "{primeira[0].strip()}").',
                )
            )
            continue
        cabecalho = normalizar_rotulo(primeira[0])
        m_testemunha = _RE_TESTEMUNHA_HEADER.match(cabecalho)
        if cabecalho in SECOES_CAMPOS or cabecalho == "informacoes da cliente":
            _processar_tabela_principal(tabela, dados, pendencias, encontrados)
        elif m_testemunha:
            indice = int(m_testemunha.group(1)) - 1
            if 0 <= indice < len(dados["testemunhas"]):
                _processar_tabela_testemunha(tabela, indice, dados, pendencias)
        else:
            pendencias.append(
                PendenciaColeta(
                    CAMPO_DOCUMENTO,
                    f'Não foi possível interpretar uma tabela do documento (seção "{primeira[0].strip()}").',
                )
            )

    # Rótulo ausente (a linha inteira não apareceu no arquivo) também vira
    # pendência — diferente de "apareceu mas veio vazio", que só falta no
    # formulário como qualquer campo em branco de preenchimento manual.
    for secao, campos in SECOES_CAMPOS.items():
        for rotulo, campo in campos.items():
            if (secao, rotulo) not in encontrados:
                pendencias.append(PendenciaColeta(campo.caminho, f'"{campo.rotulo}" não encontrado no documento.'))
    if ("informacoes sobre o servico", ROTULO_ETAPAS) not in encontrados:
        pendencias.append(PendenciaColeta("projeto.escopos", '"Etapas do serviço" não encontrado no documento.'))

    # Parcela é sempre derivada (mesma regra do formulário manual): calcula só
    # depois de ter valor_total e numero_parcelas, nunca lida do documento.
    numero_parcelas = dados["financeiro"]["numero_parcelas"]
    if numero_parcelas:
        dados["financeiro"]["valor_parcela"] = round(dados["financeiro"]["valor_total"] / numero_parcelas, 2)

    # Mesmo comportamento do preenchimento manual: cobrança usa o e-mail do
    # representante por padrão.
    dados["contratante"]["email_cobranca"] = dados["contratante"]["representante"]["email"]

    return dados, pendencias
