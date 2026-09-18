"""Editar o texto do rascunho gerado, parágrafo por parágrafo (§ Contratos, 2026-09-18).

Usa um .docx real em disco (o que `editar_docx.py` manipula) e monkeypatch
só na conversão pra PDF (dependência externa do LibreOffice, não é o que
este teste cobre).
"""

from types import SimpleNamespace

import pytest
from docx import Document

import src.use_cases.documento_contratual.editar_texto as editar_texto_mod
from src.use_cases.documento_contratual.editar_texto import (
    EditarTextoDocumentoContratualUseCase,
    GetParagrafosEditaveisUseCase,
)
from src.utils.exceptions import RegraDeNegocioError


class FakeDocumentoRepo:
    def __init__(self, documento):
        self._documento = documento

    def get_by_id(self, _id):
        return self._documento


class FakeVersaoRepo:
    def __init__(self, versao):
        self._versao = versao

    def ultima_versao_obj(self, _documento_id):
        return self._versao


def _docx_de_teste(tmp_path):
    doc = Document()
    doc.add_paragraph("Cláusula 1ª Texto original.")
    caminho = tmp_path / "rascunho.docx"
    doc.save(caminho)
    return str(caminho)


def documento(status="em_revisao_interna"):
    return SimpleNamespace(id=1, status=status)


def montar_editar(doc, versao, monkeypatch):
    monkeypatch.setattr(editar_texto_mod, "converter_docx_para_pdf", lambda *_a, **_k: None)
    uc = EditarTextoDocumentoContratualUseCase.__new__(EditarTextoDocumentoContratualUseCase)
    uc.documentos = FakeDocumentoRepo(doc)
    uc.versoes = FakeVersaoRepo(versao)
    return uc


def montar_get(doc, versao):
    uc = GetParagrafosEditaveisUseCase.__new__(GetParagrafosEditaveisUseCase)
    uc.documentos = FakeDocumentoRepo(doc)
    uc.versoes = FakeVersaoRepo(versao)
    return uc


class TestGetParagrafosEditaveis:
    def test_le_do_arquivo_da_ultima_versao(self, tmp_path):
        caminho = _docx_de_teste(tmp_path)
        versao = SimpleNamespace(docx_path=caminho)
        uc = montar_get(documento(), versao)

        paragrafos = uc.execute(1)

        assert paragrafos == [{"ref": 0, "texto": "Cláusula 1ª Texto original."}]

    def test_recusa_fora_de_status_de_edicao(self, tmp_path):
        versao = SimpleNamespace(docx_path=_docx_de_teste(tmp_path))
        uc = montar_get(documento(status="aguardando_aprovacao_cliente"), versao)

        with pytest.raises(RegraDeNegocioError, match="editar o texto"):
            uc.execute(1)

    def test_recusa_sem_versao_gerada(self):
        uc = montar_get(documento(), None)

        with pytest.raises(RegraDeNegocioError, match="Nenhum rascunho"):
            uc.execute(1)


class TestEditarTexto:
    def test_aplica_e_converte_pra_pdf(self, tmp_path, monkeypatch):
        caminho = _docx_de_teste(tmp_path)
        versao = SimpleNamespace(docx_path=caminho)
        uc = montar_editar(documento(), versao, monkeypatch)

        resultado = uc.execute(1, {0: "Cláusula 1ª Texto revisado."})

        assert resultado == {"alterados": 1}
        assert Document(caminho).paragraphs[0].text == "Cláusula 1ª Texto revisado."

    def test_recusa_fora_de_status_de_edicao(self, tmp_path, monkeypatch):
        versao = SimpleNamespace(docx_path=_docx_de_teste(tmp_path))
        uc = montar_editar(documento(status="assinado_e_arquivado"), versao, monkeypatch)

        with pytest.raises(RegraDeNegocioError, match="editar o texto"):
            uc.execute(1, {0: "x"})

    def test_recusa_sem_edicoes(self, tmp_path, monkeypatch):
        versao = SimpleNamespace(docx_path=_docx_de_teste(tmp_path))
        uc = montar_editar(documento(), versao, monkeypatch)

        with pytest.raises(RegraDeNegocioError, match="Nenhuma edição"):
            uc.execute(1, {})
