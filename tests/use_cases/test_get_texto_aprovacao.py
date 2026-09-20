"""Texto do documento em blocos pra tela pública citar trechos (§ Contratos, 2026-09-20)."""

import io
from types import SimpleNamespace

import pytest
from docx import Document

from src.use_cases.documento_contratual.get_texto_aprovacao import (
    GetTextoAprovacaoPorTokenUseCase,
)
from src.utils.exceptions import RegraDeNegocioError


class FakeRepo:
    def __init__(self, registro=None):
        self._registro = registro

    def get_by_token(self, _token):
        return self._registro

    def get_by_id(self, _id):
        return self._registro


def _docx_bytes():
    doc = Document()
    doc.add_paragraph("Cláusula 1ª O objeto deste contrato é a prestação de serviços.")
    doc.add_paragraph("Cláusula 2ª O prazo é de 30 dias úteis.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def montar(token, versao):
    uc = GetTextoAprovacaoPorTokenUseCase.__new__(GetTextoAprovacaoPorTokenUseCase)
    uc.tokens = FakeRepo(token)
    uc.versoes = FakeRepo(versao)
    return uc


def test_devolve_os_paragrafos_do_texto():
    tok = SimpleNamespace(versao_id=9)
    versao = SimpleNamespace(docx_conteudo=_docx_bytes())
    uc = montar(tok, versao)

    paragrafos = uc.execute("qualquer-token")

    assert paragrafos == [
        "Cláusula 1ª O objeto deste contrato é a prestação de serviços.",
        "Cláusula 2ª O prazo é de 30 dias úteis.",
    ]


def test_recusa_token_inexistente():
    uc = montar(None, None)

    with pytest.raises(RegraDeNegocioError, match="Link inválido"):
        uc.execute("x")


def test_recusa_versao_sem_conteudo():
    tok = SimpleNamespace(versao_id=9)
    versao = SimpleNamespace(docx_conteudo=None)
    uc = montar(tok, versao)

    with pytest.raises(RegraDeNegocioError, match="Link inválido"):
        uc.execute("x")
