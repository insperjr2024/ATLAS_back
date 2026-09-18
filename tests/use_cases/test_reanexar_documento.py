"""Reanexar um .docx editado fora da plataforma (§ Contratos, 2026-09-18)."""

from types import SimpleNamespace

import pytest
from docx import Document

import src.use_cases.documento_contratual.reanexar_documento as reanexar_mod
from src.use_cases.documento_contratual.reanexar_documento import (
    ReanexarDocumentoContratualUseCase,
)
from src.utils.exceptions import RegraDeNegocioError


class FakeDocumentoRepo:
    def __init__(self, documento):
        self._documento = documento

    def get_by_id(self, _id):
        return self._documento

    def update(self, _id, **kwargs):
        for chave, valor in kwargs.items():
            setattr(self._documento, chave, valor)
        return self._documento


class FakeVersaoRepo:
    def __init__(self):
        self.criada = None

    def ultima_versao(self, _documento_id):
        return 1

    def create(self, **kwargs):
        self.criada = SimpleNamespace(**kwargs)
        return self.criada


def _docx_bytes():
    import io

    doc = Document()
    doc.add_paragraph("Conteúdo editado fora da plataforma.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def documento(status="em_revisao_interna"):
    return SimpleNamespace(id=1, tipo="contrato", status=status, projeto_id=7)


def montar(doc, tmp_path, monkeypatch):
    monkeypatch.setattr(reanexar_mod, "GERADOS_DIR", str(tmp_path))
    monkeypatch.setattr(reanexar_mod, "converter_docx_para_pdf", lambda docx_path: docx_path[:-5] + ".pdf")
    uc = ReanexarDocumentoContratualUseCase.__new__(ReanexarDocumentoContratualUseCase)
    uc.documentos = FakeDocumentoRepo(doc)
    uc.versoes = FakeVersaoRepo()
    return uc


def test_cria_nova_versao_a_partir_do_arquivo_enviado(tmp_path, monkeypatch):
    doc = documento()
    uc = montar(doc, tmp_path, monkeypatch)

    versao = uc.execute(1, _docx_bytes())

    assert versao.versao == 2
    assert versao.status_arquivo == "rascunho"
    assert Document(versao.docx_path).paragraphs[0].text == "Conteúdo editado fora da plataforma."
    assert doc.status == "em_revisao_interna"


def test_recusa_sem_arquivo(tmp_path, monkeypatch):
    uc = montar(documento(), tmp_path, monkeypatch)

    with pytest.raises(RegraDeNegocioError, match="Envie um arquivo"):
        uc.execute(1, b"")


def test_recusa_fora_de_status_de_edicao(tmp_path, monkeypatch):
    uc = montar(documento(status="aprovado_pelo_cliente"), tmp_path, monkeypatch)

    with pytest.raises(RegraDeNegocioError, match="anexar um novo rascunho"):
        uc.execute(1, _docx_bytes())


def test_recusa_arquivo_que_nao_e_docx_valido(tmp_path, monkeypatch):
    uc = montar(documento(), tmp_path, monkeypatch)

    with pytest.raises(RegraDeNegocioError, match="não é um .docx válido"):
        uc.execute(1, b"isto nao e um docx")
