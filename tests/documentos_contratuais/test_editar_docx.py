"""Edição manual do texto de um .docx real (§ Contratos, 2026-09-18).

Sem mock: exercita `python-docx` de verdade contra um arquivo temporário —
é o único jeito de garantir que a formatação (negrito do rótulo da cláusula)
sobrevive à edição.
"""

import pytest
from docx import Document

from src.documentos_contratuais.editar_docx import aplicar_edicoes, extrair_paragrafos_editaveis


def _docx_de_teste(tmp_path):
    doc = Document()

    # Parágrafo com dois runs: rótulo em negrito + corpo — o caso que a
    # edição precisa preservar.
    p = doc.add_paragraph()
    rotulo = p.add_run("Cláusula 1ª ")
    rotulo.bold = True
    p.add_run("O objeto deste contrato é a prestação de serviços.")

    doc.add_paragraph("Parágrafo simples, um run só.")
    doc.add_paragraph("")  # separador vazio, não deve aparecer como editável
    doc.add_paragraph("INSPER JR\t\t\tCONTRATANTE")  # linha de assinatura, tabulada

    caminho = tmp_path / "teste.docx"
    doc.save(caminho)
    return str(caminho)


def test_extrai_so_os_paragrafos_editaveis(tmp_path):
    caminho = _docx_de_teste(tmp_path)

    editaveis = extrair_paragrafos_editaveis(caminho)

    assert [e["texto"] for e in editaveis] == [
        "Cláusula 1ª O objeto deste contrato é a prestação de serviços.",
        "Parágrafo simples, um run só.",
    ]
    # `ref` é a posição REAL em doc.paragraphs, não reindexada — o vazio e a
    # linha de assinatura (índices 2 e 3) ficam de fora, mas não deslocam
    # os refs dos que sobraram.
    assert [e["ref"] for e in editaveis] == [0, 1]


def test_aplicar_edicoes_preserva_negrito_do_primeiro_run(tmp_path):
    caminho = _docx_de_teste(tmp_path)

    alterados = aplicar_edicoes(caminho, {0: "Cláusula 1ª Texto reescrito pelo Jurídico."})

    assert alterados == 1
    doc = Document(caminho)
    paragrafo = doc.paragraphs[0]
    assert paragrafo.text == "Cláusula 1ª Texto reescrito pelo Jurídico."
    # O primeiro run levou o texto todo; os demais foram esvaziados — mas o
    # negrito do run continua marcado nele.
    assert paragrafo.runs[0].bold is True


def test_nao_toca_paragrafo_de_assinatura(tmp_path):
    caminho = _docx_de_teste(tmp_path)

    editaveis = extrair_paragrafos_editaveis(caminho)
    refs = [e["ref"] for e in editaveis]

    assert 3 not in refs  # a linha "INSPER JR\t\t\tCONTRATANTE"


def test_recusa_ref_fora_do_alcance(tmp_path):
    caminho = _docx_de_teste(tmp_path)

    with pytest.raises(ValueError, match="mudou desde"):
        aplicar_edicoes(caminho, {99: "x"})


def test_nao_altera_quando_texto_e_igual(tmp_path):
    caminho = _docx_de_teste(tmp_path)

    alterados = aplicar_edicoes(caminho, {1: "Parágrafo simples, um run só."})

    assert alterados == 0
